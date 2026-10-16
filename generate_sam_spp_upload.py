"""
generate_sam_spp_upload.py

Reads the "SPP Export" sheet from a CR SAM NextGen Goal File (e.g.
"MO_SAM_NextGen_<Month>_<Year>_SPP_Export_v#.xlsx") and generates LPM-format
CSV files ready for upload to LiquidPerform. Fork of generate_lpm_upload.py
(the standard SPP generator) — the CR SAM template's "SPP Export" sheet is
already laid out in the exact same column order as the standardized
Tracking Table (confirmed via the sheet's own note in cell AP2: "Columns
A:AA are the SPP Tracking Table in its own order"), so this reuses the SPP
generator's grouping, Tracker/PTG building, and collection-lookup logic
unchanged. Only the row-reading step differs:

  - Source sheet is "SPP Export", not "Tracking Table".
  - Columns 23/24, which the standard Tracking Table calls "Min Cases" /
    "Min Facings" (unused — the generator doesn't do SET, and SET isn't
    approved for SPPs), are repurposed in the CR SAM template as
    "Distribution Qualifier Qty" / "Distribution Qualifier Unit" (e.g. qty=3,
    unit="bottle"/"case"). Confirmed: these only populate on
    ACS/NACS/POD/NPOD (distribution) objectives. Qty feeds achievement_min
    (taking priority over the legacy col-21 "Qualifier" value, which SAM
    exports leave blank for these objective types); Unit overrides the
    Measure-derived goal_uom for that row ("bottle" -> "Bottles",
    "case"/"cases" -> "Cases").

Usage:
    python generate_sam_spp_upload.py <sam_goal_file.xlsx> [output_csv] [collection_report.xlsx]

Open assumptions (confirm/adjust before relying on this in production):
  - Only verified against the MO November 2026 export. Other CR SAM states
    are assumed to share the same "SPP Export" column layout but have not
    been checked yet.
  - Distribution Qualifier Unit values other than bottle/case (singular or
    plural) are not recognized and fall back to the standard Measure-based
    UOM default (blank unit text is simply ignored, same as today).
"""

import os
import sys

try:
    import openpyxl
except ImportError:
    print("ERROR: openpyxl is required.  Install with: pip install openpyxl", file=sys.stderr)
    sys.exit(1)

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import generate_lpm_upload as spp  # reuse CSV schema, constants, and most of the pipeline

CSV_COLUMNS = spp.CSV_COLUMNS

SOURCE_SHEET_NAME = "SPP Export"

# Distribution Qualifier Unit (col 24) -> LPM goal_uom
DIST_QUALIFIER_UOM_MAP = {
    "bottle":  "Bottles",
    "bottles": "Bottles",
    "case":    "Cases",
    "cases":   "Cases",
}


def load_sam_spp_export(wb_path):
    """
    Read the "SPP Export" sheet and return (records, skipped_rows, header_row),
    in the same shape as generate_lpm_upload.load_tracking_table().
    """
    wb_formula = openpyxl.load_workbook(wb_path, read_only=True, data_only=False)
    wb = openpyxl.load_workbook(wb_path, read_only=True, data_only=True)

    if SOURCE_SHEET_NAME not in wb.sheetnames:
        raise ValueError(
            f"No '{SOURCE_SHEET_NAME}' sheet found in {wb_path}.\n"
            f"Available sheets: {wb.sheetnames}"
        )

    ws_formula = wb_formula[SOURCE_SHEET_NAME]
    ws = wb[SOURCE_SHEET_NAME]

    # Read_only sheets only support fast sequential access (iter_rows) — random
    # .cell(r, c) access there re-scans from the top of the sheet every call,
    # which is O(row number) per call and O(n^2) over a full pass. For a
    # large multi-sheet CR SAM workbook that turns a sub-second read into a
    # multi-minute hang, so both the cached and formula views are walked in
    # lockstep instead.
    formula_rows = ws_formula.iter_rows(min_row=1, max_col=27, values_only=True)
    data_rows    = ws.iter_rows(min_row=1, max_col=27, values_only=True)
    header_row = list(next(data_rows))
    next(formula_rows)

    records = []
    skipped = []

    def make_skip(row_num, reason, raw):
        return {"row_num": row_num, "reason": reason, "raw_row": raw}

    for r, (row_formula, row_data) in enumerate(zip(formula_rows, data_rows), start=2):
        def cv(c):
            idx = c - 1
            if c in (5, 6, 7, 8):
                return row_data[idx]
            v = row_formula[idx]
            if isinstance(v, str) and v.startswith("="):
                return row_data[idx]
            return v

        goal_group  = spp.safe_str(cv(1))
        spp_tier    = spp.safe_str(cv(2))
        goal_bucket = spp.safe_str(cv(3))
        obj_type    = spp.safe_str(cv(4)).rstrip(" -")
        start_raw        = cv(5)
        end_raw          = cv(6)
        basis_start_raw  = cv(7)
        basis_end_raw    = cv(8)
        unsold_prd  = spp.safe_str(cv(9))
        measure     = spp.safe_str(cv(10))
        mkt_seg     = cv(11)
        if isinstance(mkt_seg, str) and mkt_seg.strip().upper() == "FLAT":
            mkt_seg = 0
        elif isinstance(mkt_seg, float) and -1 < mkt_seg < 1 and mkt_seg != 0:
            mkt_seg = round(mkt_seg * 100, 1 if mkt_seg < 0 else 2)
        elif isinstance(mkt_seg, (int, float)):
            mkt_seg = round(mkt_seg, 1 if mkt_seg < 0 else 2)
        ptg_name_v  = cv(12)
        level_detail = spp.safe_str(cv(13))
        supplier    = cv(14)
        selection   = cv(15)
        applicable_premise = spp.safe_str(cv(18)).upper()
        goal_distribution = spp.safe_str(cv(20))
        qualifier_legacy   = cv(21)   # "Qualifier" — blank on SAM distribution rows
        min_goal_per_rep   = cv(22)
        dist_qual_qty      = cv(23)   # "Distribution Qualifier Qty", e.g. 3
        dist_qual_unit     = spp.safe_str(cv(24)).lower()  # "Distribution Qualifier Unit", e.g. "bottle"
        pod_attr    = cv(25)
        category    = spp.safe_str(cv(27))

        raw_row = list(row_data)

        if not goal_group and ptg_name_v is None and selection is None:
            continue

        if measure in spp.SKIP_MEASURES:
            continue

        if goal_bucket == "Select:" and not obj_type:
            continue

        if obj_type in spp.SKIP_OBJECTIVE_TYPES:
            skipped.append(make_skip(r, f"skipped objective type: {obj_type!r}", raw_row))
            continue

        goal_type = spp.OBJECTIVE_TYPE_MAP.get(obj_type)
        if not goal_type:
            skipped.append(make_skip(r, f"unknown objective type: {obj_type!r}", raw_row))
            continue

        _unsold_present = unsold_prd and unsold_prd.strip().upper() != "NONE"
        _unsold_none    = unsold_prd and unsold_prd.strip().upper() == "NONE"
        if goal_type == "DistributionPODs" and _unsold_present:
            goal_type = "DistributionNewPODs"
        elif goal_type == "DistributionACS" and _unsold_present:
            goal_type = "DistributionNewACS"
        elif goal_type == "DistributionNewPODs" and _unsold_none:
            goal_type = "DistributionPODs"
        elif goal_type == "DistributionNewACS" and _unsold_none:
            goal_type = "DistributionACS"

        name = spp.ptg_name_from_row(ptg_name_v, selection, supplier)
        if not name and level_detail == "Total":
            name = f"Total {obj_type}"
        if not name:
            skipped.append(make_skip(r, "no PTG name, selection, or supplier", raw_row))
            continue

        if goal_type in {"VolumeRevenue", "VolumeRevenueDREV", "VolumeGP", "VolumeCombos", "VolumePoints"}:
            goal_uom = ""
        else:
            goal_uom = spp.UOM_MAP.get(measure, "")
            if not goal_uom and goal_type in spp.DISTRIBUTION_TYPES:
                goal_uom = "Cases"

        # Distribution Qualifier Unit overrides the Measure-derived UOM for
        # distribution-type goals when recognized.
        if goal_type in spp.DISTRIBUTION_TYPES and dist_qual_unit in DIST_QUALIFIER_UOM_MAP:
            goal_uom = DIST_QUALIFIER_UOM_MAP[dist_qual_unit]

        # Distribution Qualifier Qty takes priority over the legacy Qualifier
        # column for distribution-type goals; falls back to col 21 otherwise.
        if goal_type in spp.DISTRIBUTION_TYPES and spp._numeric_str(dist_qual_qty):
            qualifier = dist_qual_qty
        else:
            qualifier = qualifier_legacy

        start_yyyymm       = spp.to_yyyymm(start_raw)
        end_yyyymm         = spp.to_yyyymm(end_raw)
        basis_start_yyyymm = spp.to_yyyymm(basis_start_raw)
        basis_end_yyyymm   = spp.to_yyyymm(basis_end_raw)

        pod_attr_clean = spp.normalize_pod_attribute(pod_attr)
        if pod_attr_clean and pod_attr_clean not in spp.VALID_POD_ATTRIBUTES:
            skipped.append(make_skip(r, f"invalid POD attribute: {pod_attr_clean!r}", raw_row))
            continue

        records.append({
            "row_num":       r,
            "goal_group":    goal_group,
            "spp_tier":      spp_tier,
            "goal_type":     goal_type,
            "goal_uom":      goal_uom,
            "program_class": spp.PROGRAM_CLASS_MAP.get(spp_tier, ""),
            "start_yyyymm":        start_yyyymm,
            "end_yyyymm":          end_yyyymm,
            "basis_start_yyyymm":  basis_start_yyyymm,
            "basis_end_yyyymm":    basis_end_yyyymm,
            "unsold_prd":          unsold_prd,
            "supplier":      spp.safe_str(supplier),
            "ptg_name":      name,
            "mkt_seg_goal":       spp._numeric_str(mkt_seg),
            "pod_attribute":      pod_attr_clean,
            "min_goal_per_rep":   spp._numeric_str(min_goal_per_rep),
            "qualifier":          spp._numeric_str(qualifier),
            "goal_distribution":  goal_distribution,
            "level_detail":       level_detail,
            "selection":          spp.safe_str(selection),
            "category":           category,
            "premise":            applicable_premise,
        })

    return records, skipped, header_row


def main():
    if len(sys.argv) < 2:
        print(
            "Usage: python generate_sam_spp_upload.py <sam_goal_file.xlsx> [output.csv] [collection_report.xlsx]",
            file=sys.stderr,
        )
        sys.exit(1)

    input_path = sys.argv[1]
    if not os.path.isfile(input_path):
        print(f"ERROR: File not found: {input_path}", file=sys.stderr)
        sys.exit(1)

    if len(sys.argv) >= 3:
        output_path = sys.argv[2]
    else:
        base = os.path.splitext(os.path.abspath(input_path))[0]
        output_path = base + "_lpm_upload.csv"

    collection_path = sys.argv[3] if len(sys.argv) >= 4 else None
    if not collection_path:
        folder = os.path.dirname(os.path.abspath(input_path))
        import glob as _glob
        candidates = _glob.glob(os.path.join(folder, "*Collection*Report*.xlsx"))
        if candidates:
            collection_path = candidates[0]

    base = os.path.splitext(output_path)[0]
    skipped_path = base + "_skipped.csv"

    fiscal_path = os.path.join(os.path.dirname(os.path.abspath(__file__)), "Supplier Fiscal Start Month.xlsx")
    if not os.path.isfile(fiscal_path):
        fiscal_path = os.path.join(os.path.dirname(os.path.abspath(input_path)), "Supplier Fiscal Start Month.xlsx")
    if os.path.isfile(fiscal_path):
        spp.FISCAL_LOOKUP = spp.load_fiscal_lookup(fiscal_path)
        print(f"Fiscal lookup loaded: {len(spp.FISCAL_LOOKUP)} suppliers  ({os.path.basename(fiscal_path)})")
    else:
        print("Fiscal lookup not found — FYTD unsold periods will fall back to CYTD.")

    wb_for_fiscal = openpyxl.load_workbook(input_path, read_only=True, data_only=True)
    spp.INTERNAL_FISCAL_LOOKUP = spp.load_internal_fiscal_lookup(wb_for_fiscal)

    state_from_sheet = ""
    if SOURCE_SHEET_NAME in wb_for_fiscal.sheetnames:
        tt = wb_for_fiscal[SOURCE_SHEET_NAME]
        for (gg,) in tt.iter_rows(min_row=2, max_col=1, values_only=True):
            if gg:
                state_from_sheet = spp.derive_state_from_goal_group(str(gg))
                if state_from_sheet:
                    break
    wb_for_fiscal.close()

    state = state_from_sheet or spp.derive_state_from_filename(input_path)
    spp.CURRENT_STATE = state
    if collection_path and os.path.isfile(collection_path):
        spp.COLLECTION_LOOKUP = spp.load_collection_lookup(collection_path, state)
        print(f"State: {state}  |  Collection IDs loaded: {len(spp.COLLECTION_LOOKUP)}  ({os.path.basename(collection_path)})")
        for name, cid in sorted(spp.COLLECTION_LOOKUP.items()):
            print(f"  {name} -> {cid}")
    else:
        print(f"State: {state}  |  No collection report found — salesforce_collection_ids will be blank.")

    print(f"Internal FISCAL sheet entries: {len(spp.INTERNAL_FISCAL_LOOKUP)}")

    print(f"Reading: {input_path}")
    records, skipped, header_row = load_sam_spp_export(input_path)

    if not records:
        print(f"ERROR: No valid records found in '{SOURCE_SHEET_NAME}' after filtering.", file=sys.stderr)
        sys.exit(1)

    order, groups = spp.group_records(records)
    total_trackers, total_ptgs = spp.generate_output(order, groups, output_path)
    duplicate_names = spp.find_duplicate_ptg_names(order, groups)

    if skipped:
        spp.write_skipped_csv(skipped_path, skipped, header_row)

    spp.print_summary(output_path, skipped_path, total_trackers, total_ptgs, skipped, duplicate_names)


if __name__ == "__main__":
    main()
