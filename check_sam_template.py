"""
check_sam_template.py

Pre-flight checker for CR SAM NextGen Goal File .xlsx files. Scans the
"SPP Export" sheet for data-entry problems that generate_sam_spp_upload.py
either silently coerces, silently skips, or doesn't catch at all, and
reports them as flags — this never blocks CSV generation, it just surfaces
things worth a second look before the state's submission goes through.

Fork of check_spp_template.py (the standard SPP checker) — the CR SAM
"SPP Export" sheet shares the same column 1-27 layout as the standardized
Tracking Table, except columns 23/24 are repurposed as "Distribution
Qualifier Qty" / "Distribution Qualifier Unit" instead of "Min Cases" /
"Min Facings / Mentions" (see generate_sam_spp_upload.py). Reuses
generate_lpm_upload's helpers the same way the SAM generator does. The
workbook also tends to be a large multi-sheet file (DM/PACE/Forecast tabs,
monthly tabs, etc.), so rows are read with read_only + lockstep
formula/cached iteration like generate_sam_spp_upload.load_sam_spp_export,
rather than the random-access .cell(r, c) style check_spp_template.py uses
for the much smaller single-purpose Goal Builder .xlsm.

Usage:
    python check_sam_template.py <sam_goal_file.xlsx>
"""

import csv
import sys

try:
    import openpyxl
except ImportError:
    print("ERROR: openpyxl is required.  Install with: pip install openpyxl", file=sys.stderr)
    sys.exit(1)

import generate_lpm_upload as gen
import generate_sam_spp_upload as samgen

SOURCE_SHEET_NAME = samgen.SOURCE_SHEET_NAME  # "SPP Export"

# Canonical SPP Export header (columns 1-27), whitespace-normalized. Same as
# check_spp_template.EXPECTED_HEADERS except columns 23/24, which the CR SAM
# template repurposes for the Distribution Qualifier Qty/Unit fields.
EXPECTED_HEADERS = [
    "Goal Group",
    "SPP Tier",
    "Goal Bucket:",
    "Objective Type",
    "Program Posting Start",
    "Program Posting End",
    "Basis Period Start",
    "Basis Period End",
    "Unsold Prd",
    "Measure",
    "Market Segment Goal",
    "PTG Name (Brand/Product Info ONLY):",
    "Level of Detail",
    "Supplier",
    "Selection (select all that apply)",
    "Customer Exclusions",
    "Basis Item (if different than selection)",
    "Applicable Premise",
    "Size(s)",
    "Goal Distribution",
    "Qualifier",
    "Min Goal per Rep",
    "Distribution Qualifier Qty",
    "Distribution Qualifier Unit",
    "POD Attribute",
    "Notes/Exclusions",
    "Category",
]


def _norm_header(v):
    """Collapse whitespace/newlines so '<br>'-wrapped headers still compare cleanly."""
    if v is None:
        return ""
    return " ".join(str(v).split())


def _is_self_concat(s):
    """True if s is exactly two back-to-back copies of the same non-empty text,
    e.g. 'Volume (Cases)Volume (Cases)' -- a drag-fill/paste artifact."""
    if not isinstance(s, str):
        return False
    s = s.strip()
    if len(s) < 2 or len(s) % 2 != 0:
        return False
    half = len(s) // 2
    first, second = s[:half], s[half:]
    return bool(first) and first == second


def _to_float(value):
    """Best-effort numeric coercion, reusing gen._numeric_str for strings
    with trailing non-numeric text (e.g. '12 pods')."""
    if value is None:
        return None
    if isinstance(value, (int, float)):
        return float(value)
    s = gen._numeric_str(value)
    if not s:
        return None
    try:
        return float(s)
    except ValueError:
        return None


def make_flag(row_num, severity, category, message):
    return {"row_num": row_num, "severity": severity, "category": category, "message": message}


# Level of Detail values that are backed by the VSTACK master list (dependent
# dropdown source), when one is present. CR SAM exports have not been seen to
# carry a VSTACK sheet at all -- in that case this check simply never fires.
LEVELS_IN_VSTACK = {"Supplier", "Group", "SubGroup", "Super Group", "Items(s)", "Brand"}


def _build_vstack_index(wb):
    """
    VSTACK columns: A=Category (level), B=parent Supplier (blank for the flat
    Supplier level itself), C=Concat, D=Specific (display name).
    Returns (valid, reverse_same_supplier, reverse_any_supplier):
      valid: set of (category, supplier_key, specific_key)
      reverse_same_supplier: (supplier_key, specific_key) -> set of categories
      reverse_any_supplier: specific_key -> set of (category, supplier_key)
    """
    valid = set()
    reverse_same_supplier = {}
    reverse_any_supplier = {}
    if "VSTACK" not in wb.sheetnames:
        return valid, reverse_same_supplier, reverse_any_supplier

    ws = wb["VSTACK"]
    for row in ws.iter_rows(min_row=2, values_only=True):
        cat, supplier, _concat, specific = row[0], row[1], row[2], row[3]
        if cat is None or specific is None:
            continue
        supplier_key = (supplier or "").strip().upper()
        specific_key = str(specific).strip().upper()
        valid.add((cat, supplier_key, specific_key))
        reverse_same_supplier.setdefault((supplier_key, specific_key), set()).add(cat)
        reverse_any_supplier.setdefault(specific_key, set()).add((cat, supplier_key))
    return valid, reverse_same_supplier, reverse_any_supplier


def check_tracking_table(wb_path):
    """
    Return (flags, header_issues, row_count).
    flags: list of dicts (row_num=None for file-level flags).
    header_issues: list of strings describing header/column problems.
    """
    flags = []
    header_issues = []

    # read_only + lockstep formula/cached iteration: CR SAM workbooks carry
    # many other sheets (PACE/Forecast/monthly tabs, etc.) and random .cell(r, c)
    # access on a read_only sheet re-scans from the top every call, which is
    # fine for a Goal Builder .xlsm but turns into a multi-minute hang here.
    wb_formula = openpyxl.load_workbook(wb_path, read_only=True, data_only=False)
    wb = openpyxl.load_workbook(wb_path, read_only=True, data_only=True)

    if SOURCE_SHEET_NAME not in wb.sheetnames:
        header_issues.append(f"No '{SOURCE_SHEET_NAME}' sheet found in this file.")
        return flags, header_issues, 0

    ws_formula = wb_formula[SOURCE_SHEET_NAME]
    ws = wb[SOURCE_SHEET_NAME]
    vstack_valid, vstack_same_supplier, vstack_any_supplier = _build_vstack_index(wb)

    formula_rows = ws_formula.iter_rows(min_row=1, max_col=27, values_only=True)
    data_rows    = ws.iter_rows(min_row=1, max_col=27, values_only=True)
    header_row_formula = list(next(formula_rows))
    header_row = list(next(data_rows))

    # --- Header / missing-column check -------------------------------------
    actual_headers = [_norm_header(v) for v in header_row[: len(EXPECTED_HEADERS)]]
    if len(header_row) < len(EXPECTED_HEADERS):
        header_issues.append(
            f"'{SOURCE_SHEET_NAME}' has only {len(header_row)} column(s); expected {len(EXPECTED_HEADERS)}. "
            "Columns are likely missing from the template."
        )
    for i, expected in enumerate(EXPECTED_HEADERS):
        actual = actual_headers[i] if i < len(actual_headers) else ""
        if actual != expected:
            col_letter = openpyxl.utils.get_column_letter(i + 1)
            header_issues.append(
                f"Column {col_letter}: expected header \"{expected}\", found \"{actual or '(blank)'}\"."
            )

    end_dates_seen = {}  # yyyymm -> list of row_num
    tier_counts = {}     # goal_group -> {"Anchor": n, "Flex": n}
    row_count = 0

    for r, (row_formula, row_data) in enumerate(zip(formula_rows, data_rows), start=2):
        def cv(c):
            idx = c - 1
            if c in (5, 6, 7, 8):
                return row_data[idx]
            v = row_formula[idx]
            if isinstance(v, str) and v.startswith("="):
                return row_data[idx]
            return v

        goal_group  = gen.safe_str(cv(1))
        spp_tier    = gen.safe_str(cv(2))
        goal_bucket = gen.safe_str(cv(3))
        obj_type    = gen.safe_str(cv(4)).rstrip(" -")
        start_raw   = cv(5)
        end_raw     = cv(6)
        mkt_seg     = cv(11)
        ptg_name_v  = cv(12)
        level_detail = gen.safe_str(cv(13))
        supplier    = cv(14)
        selection   = cv(15)
        measure     = gen.safe_str(cv(10))
        goal_distribution    = gen.safe_str(cv(20))
        min_goal_per_rep_raw = cv(22)
        dist_qual_qty        = cv(23)  # "Distribution Qualifier Qty"
        dist_qual_unit       = gen.safe_str(cv(24)).lower()  # "Distribution Qualifier Unit"

        if not goal_group and ptg_name_v is None and selection is None:
            continue  # genuinely blank row — not a problem

        row_count += 1

        # --- Anchor/Flex tally (2:1 ratio check, below) ----------------------
        if spp_tier in ("Anchor", "Flex"):
            counts = tier_counts.setdefault(goal_group or "(no Goal Group)", {"Anchor": 0, "Flex": 0})
            counts[spp_tier] += 1

        # --- Blank end date -------------------------------------------------
        start_yyyymm = gen.to_yyyymm(start_raw)
        end_yyyymm   = gen.to_yyyymm(end_raw)
        if start_yyyymm and not end_yyyymm:
            flags.append(make_flag(
                r, "warning", "end_date",
                f"Program Posting Start is set ({start_raw}) but Program Posting End is blank."
            ))
        elif end_yyyymm:
            end_dates_seen.setdefault(end_yyyymm, []).append(r)

        # --- Duplicated/concatenated text ------------------------------------
        row_vals = [cv(c) for c in range(1, len(EXPECTED_HEADERS) + 1)]
        for c, val in enumerate(row_vals, start=1):
            if _is_self_concat(val):
                header = EXPECTED_HEADERS[c - 1]
                flags.append(make_flag(
                    r, "warning", "duplicated_text",
                    f"\"{header}\" looks duplicated/concatenated: {val!r}"
                ))

        # NOTE: check_spp_template.py flags Supplier being filled in when Level
        # of Detail = "Supplier" (grayed out in the Goal Builder .xlsm, so a
        # filled value is suspicious). That doesn't translate here -- CR SAM's
        # "SPP Export" sheet populates Supplier for every Supplier-level row
        # by design (confirmed: 111/111 on the MO Nov 2026 export), so that
        # check is intentionally omitted to avoid 100% false-positive noise.

        # --- Selection not valid for the stated Level of Detail ---------------
        # Cross-reference against VSTACK, the dependent-dropdown master list,
        # when the workbook has one (CR SAM exports haven't been seen to carry
        # one, so this is typically a no-op here).
        if level_detail in LEVELS_IN_VSTACK and vstack_valid:
            chosen = gen.safe_str(selection)
            if chosen:
                supplier_key = "" if level_detail == "Supplier" else gen.safe_str(supplier).upper()

                def _is_valid(text):
                    return (level_detail, supplier_key, text.upper()) in vstack_valid

                if not _is_valid(chosen):
                    parts = [p.strip() for p in chosen.split(",") if p.strip()]
                    if len(parts) > 1 and all(_is_valid(p) for p in parts):
                        to_check = []
                    else:
                        to_check = [chosen]

                    for item in to_check:
                        item_key = item.upper()
                        same_supplier = vstack_same_supplier.get((supplier_key, item_key))
                        anywhere = vstack_any_supplier.get(item_key)
                        if same_supplier or anywhere:
                            actually = same_supplier or anywhere
                            flags.append(make_flag(
                                r, "warning", "selection_not_valid_for_level",
                                f"Selection {item!r} isn't valid as '{level_detail}' "
                                f"(Supplier {supplier_key or '(none)'!r}) — it actually exists under "
                                f"{sorted(actually)}. Likely a Level of Detail or Supplier change that "
                                f"wasn't followed by re-picking Selection."
                            ))
                        else:
                            flags.append(make_flag(
                                r, "warning", "selection_not_found",
                                f"Selection {item!r} doesn't exist in the master list under "
                                "any Level of Detail or Supplier — possibly freeform text, a typo, or a "
                                "discontinued item."
                            ))

        # --- Non-numeric Market Segment Goal ---------------------------------
        if isinstance(mkt_seg, str):
            s = mkt_seg.strip()
            if s and s.upper() != "FLAT":
                numeric = gen._numeric_str(s)
                if numeric != s:
                    flags.append(make_flag(
                        r, "warning", "non_numeric_goal",
                        f"Market Segment Goal {s!r} has non-numeric text; "
                        f"only {numeric!r} will be used."
                    ))

        # --- Fixed goal looks like a summed team total, not a per-rep target --
        MAX_PLAUSIBLE_REPS = 50
        if goal_distribution == gen.FIXED_GOAL_DISTRIBUTION:
            mkt_val = _to_float(mkt_seg)
            min_val = _to_float(min_goal_per_rep_raw)

            if mkt_val is None and min_val:
                flags.append(make_flag(
                    r, "info", "fixed_goal_from_min_per_rep",
                    f"Market Segment Goal is blank on this 'Fixed Goal per Salesperson' row, so "
                    f"Min Goal per Rep ({min_val:g}) will be used as the Fixed Goal instead."
                ))
            elif mkt_val is not None and min_val and min_val > 1 and mkt_val != min_val:
                ratio = mkt_val / min_val
                nearest = round(ratio)
                if 1 < nearest <= MAX_PLAUSIBLE_REPS and abs(ratio - nearest) < 1e-6:
                    flags.append(make_flag(
                        r, "warning", "fixed_goal_looks_summed",
                        f"Goal Distribution is 'Fixed Goal per Salesperson' with Market Segment "
                        f"Goal = {mkt_val:g}, but Min Goal per Rep = {min_val:g}. A Fixed goal is "
                        f"each rep's own target, not a team total — {mkt_val:g} looks like "
                        f"{min_val:g} x {nearest} reps summed together. This should probably be a "
                        f"Fixed Goal of {min_val:g}, not {mkt_val:g} (or switch Goal Distribution to "
                        f"'Even Goal' if {mkt_val:g} really is meant to be split across reps)."
                    ))
                else:
                    flags.append(make_flag(
                        r, "info", "fixed_goal_has_min_per_rep",
                        f"Goal Distribution is 'Fixed Goal per Salesperson', which doesn't use Min "
                        f"Goal per Rep (the Fixed Goal of {mkt_val:g} is already each rep's minimum) "
                        f"— but Min Goal per Rep is filled in as {min_val:g}. Worth confirming this "
                        f"wasn't meant to be the actual Fixed Goal."
                    ))

        # --- Distribution Qualifier Unit not recognized -----------------------
        # generate_sam_spp_upload.py only overrides goal_uom for bottle/case
        # (singular or plural); any other unit text is silently ignored and
        # the row falls back to the Measure-derived default.
        if dist_qual_unit and dist_qual_unit not in samgen.DIST_QUALIFIER_UOM_MAP:
            flags.append(make_flag(
                r, "warning", "unrecognized_dist_qualifier_unit",
                f"Distribution Qualifier Unit {dist_qual_unit!r} isn't recognized (only "
                "bottle/case are) — it will be silently ignored and the Measure-based UOM "
                "default will be used instead."
            ))

        # --- Distribution Qualifier Qty filled in without a recognized goal ---
        # Qty only feeds achievement_min on distribution-type rows; on a
        # non-distribution row it's simply never read.
        obj_type_for_goal = obj_type
        goal_type_guess = gen.OBJECTIVE_TYPE_MAP.get(obj_type_for_goal)
        if (gen._numeric_str(dist_qual_qty) and goal_type_guess
                and goal_type_guess not in gen.DISTRIBUTION_TYPES):
            flags.append(make_flag(
                r, "info", "dist_qualifier_qty_ignored",
                f"Distribution Qualifier Qty is filled in ({dist_qual_qty!r}) but Objective Type "
                f"'{obj_type}' isn't a distribution goal — this value will be ignored."
            ))

        # --- Silent skips (surfaced, not blocked) -----------------------------
        if goal_bucket == "Select:" and not obj_type:
            flags.append(make_flag(
                r, "info", "silent_skip",
                "Goal Bucket is still 'Select:' with no Objective Type chosen — "
                "this row will be dropped with no record in the skipped-rows CSV."
            ))
        elif measure in gen.SKIP_MEASURES:
            flags.append(make_flag(
                r, "info", "silent_skip",
                f"Measure '{measure}' is intentionally excluded (Digital) — "
                "this row will be dropped with no record in the skipped-rows CSV."
            ))
        elif obj_type and not goal_bucket == "Select:" and obj_type not in gen.SKIP_OBJECTIVE_TYPES:
            if not gen.OBJECTIVE_TYPE_MAP.get(obj_type):
                flags.append(make_flag(
                    r, "warning", "unknown_objective_type",
                    f"Objective Type '{obj_type}' isn't recognized — this row will be skipped."
                ))
        elif not obj_type and goal_bucket != "Select:":
            flags.append(make_flag(
                r, "warning", "blank_objective_type",
                "Objective Type is blank — this row will be skipped as unrecognized."
            ))

    # --- Anchor/Flex 2:1 ratio, per Goal Group (tracker) --------------------
    for goal_group, counts in tier_counts.items():
        anchor_n, flex_n = counts["Anchor"], counts["Flex"]
        total = anchor_n + flex_n
        if total == 0:
            continue
        expected_anchor = round(total * 2 / 3)
        expected_flex = total - expected_anchor
        if (anchor_n, flex_n) != (expected_anchor, expected_flex):
            flags.append(make_flag(
                None, "warning", "anchor_flex_ratio",
                f"'{goal_group}': {total} goal(s) should split 2:1 Anchor/Flex "
                f"({expected_anchor} Anchor / {expected_flex} Flex), but found "
                f"{anchor_n} Anchor / {flex_n} Flex."
            ))

    # --- File-level: inconsistent end dates ---------------------------------
    if len(end_dates_seen) > 1:
        breakdown = ", ".join(
            f"{ym} (row(s) {rows})" for ym, rows in sorted(end_dates_seen.items())
        )
        flags.append(make_flag(
            None, "warning", "end_date_inconsistent",
            f"Multiple different Program Posting End months found across the file: {breakdown}. "
            "Confirm this is intentional (multiple distinct programs) and not a stale/missed update."
        ))

    wb_formula.close()
    wb.close()
    return flags, header_issues, row_count


def write_report_csv(csv_path, flags, header_issues, source_filename=""):
    """
    Write the full checker report (header issues + all flags) to a CSV,
    sorted so file-level items and the sorted-by-row flags are both easy
    to scan. Suitable for emailing to a state or filing alongside the
    submission for a paper trail.
    """
    with open(csv_path, "w", newline="", encoding="utf-8-sig") as f:
        writer = csv.writer(f, quoting=csv.QUOTE_ALL)
        writer.writerow(["source_file", "row", "severity", "category", "message"])
        for issue in header_issues:
            writer.writerow([source_filename, "", "error", "header", issue])
        for flag in sorted(flags, key=lambda fl: (fl["row_num"] is None, fl["row_num"] or 0)):
            row_label = flag["row_num"] if flag["row_num"] else "File-level"
            writer.writerow([source_filename, row_label, flag["severity"], flag["category"], flag["message"]])


def main():
    if len(sys.argv) < 2:
        print("Usage: python check_sam_template.py <sam_goal_file.xlsx>", file=sys.stderr)
        sys.exit(1)

    flags, header_issues, row_count = check_tracking_table(sys.argv[1])

    print(f"Checked {row_count} row(s).\n")

    if header_issues:
        print("HEADER ISSUES:")
        for issue in header_issues:
            print(f"  - {issue}")
        print()

    if not flags:
        print("No issues found.")
        return

    for f in flags:
        where = f"Row {f['row_num']}" if f["row_num"] else "File-level"
        print(f"[{f['severity'].upper()}] {where} ({f['category']}): {f['message']}")


if __name__ == "__main__":
    main()
